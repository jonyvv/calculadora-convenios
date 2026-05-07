from io import BytesIO

from openpyxl import load_workbook

from app.infrastructure.parsers.base import DocumentParser


class ExcelParser(DocumentParser):
    def parse(self, content: bytes) -> str:
        wb = load_workbook(BytesIO(content), data_only=True)
        rows: list[str] = []
        for sheet in wb.worksheets:
            rows.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                rows.append(" | ".join("" if cell is None else str(cell) for cell in row))
        return "\n".join(rows)
